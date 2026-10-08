"""Public judge demo security tests: no real files, network, or OpenAI token."""
from __future__ import annotations

import io

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from src.api.deploy_security import DemoAccessGateway
from src.api import public_interactive as demo


@pytest.fixture(autouse=True)
def interactive_env(monkeypatch):
    monkeypatch.setenv("RISKPILOT_PUBLIC_INTERACTIVE", "1")


PASSWORD = "only-for-private-sandbox-12345"
FAKE = ("event_id,date,amount,direction,category,source_type,status,description,"
        "source_reference,due_date,expected_cash_date\n"
        "h1,2026-10-08,100,INFLOW,customer receipts,COMMITTED,ACTIVE,"
        "Fictional receipt,r1,,\n"
        "h2,2026-10-09,75,OUTFLOW,variable operating costs,MODELLED,ACTIVE,"
        "Fictional cost,r2,,\n").encode()


def mock_app(interactive=True):
    app = FastAPI()
    sessions = {}
    history = {}

    @app.get("/")
    def home():
        return {"app": "RiskPilot"}

    @app.get("/api/health")
    def health():
        return {"status": "ok"}

    @app.get("/api/scenarios")
    def scenarios():
        return ["fictional"]

    @app.post("/api/customer/import")
    def upload():
        session = "fake-cash-session"
        sessions[session] = demo.current_public_owner()
        return {"customer_session_id": session}

    @app.get("/api/command-center/fictional")
    def customer(customer_session_id: str | None = None):
        owner = demo.current_public_owner()
        if customer_session_id and sessions.get(customer_session_id) != owner:
            raise HTTPException(404, "Session not found in this browser")
        return {"customer": customer_session_id or "synthetic"}

    @app.post("/api/forecast/history/save")
    def save():
        owner = demo.current_public_owner()
        k = demo.history_storage_key("shared-history-key-by-request")
        history[k] = f"saved by {owner}"
        return {"saved": True}

    @app.get("/api/forecast/history")
    def get_history():
        return {"history": history.get(demo.history_storage_key("shared-history-key-by-request"), "empty")}

    @app.post("/api/customer/mapping/preview")
    def preview():
        return {"columns": ["date", "amount"]}

    @app.post("/api/customer/multi/import")
    def multi():
        return {"status": "synthetic imported"}

    @app.post("/api/agent")
    def agent():
        return {"answer": "fake"}

    app.add_middleware(DemoAccessGateway, password=PASSWORD, interactive=interactive)
    return app


def client(app):
    return TestClient(app, base_url="https://testserver")


def test_home_and_key_paths_are_public():
    with client(mock_app()) as c:
        assert c.get("/").status_code == 200
        assert c.get("/api/scenarios").status_code == 200
        assert c.post("/api/customer/mapping/preview").status_code == 200
        assert c.post("/api/customer/multi/import").status_code == 200
        assert c.post("/api/forecast/history/save").status_code == 200
        assert c.get("/api/forecast/history").status_code == 200
        assert c.get("/api/health").status_code == 200
        assert c.cookies.get("rp_judge_session")


def test_distinct_browser_sessions_do_not_share_history_or_import_tokens():
    app = mock_app()
    with client(app) as alice, client(app) as bob:
        alice.get("/")
        bob.get("/")
        assert alice.cookies.get("rp_judge_session") != bob.cookies.get("rp_judge_session")
        token = alice.post("/api/customer/import").json()["customer_session_id"]
        assert alice.get("/api/command-center/fictional", params={"customer_session_id": token}).status_code == 200
        assert bob.get("/api/command-center/fictional", params={"customer_session_id": token}).status_code == 404
        alice.post("/api/forecast/history/save")
        assert alice.get("/api/forecast/history").json()["history"].startswith("saved by")
        assert bob.get("/api/forecast/history").json()["history"] == "empty"


def test_cookie_tampering_creates_another_isolated_visitor():
    app = mock_app()
    with client(app) as c:
        c.get("/")
        old_cookie = c.cookies.get("rp_judge_session")
        response = c.get("/", headers={"Cookie": "rp_judge_session=attacker." + "0" * 64})
        assert response.status_code == 200
        assert old_cookie not in response.headers.get("set-cookie", "")
        assert "rp_judge_session=" in response.headers.get("set-cookie", "")


def test_private_mode_still_requires_login_and_health_is_public():
    with client(mock_app(interactive=False)) as c:
        assert c.get("/api/health").status_code == 200
        assert c.get("/").status_code == 401
        assert c.get("/api/scenarios").status_code == 401


def test_api_docs_and_unknown_writes_are_blocked():
    with client(mock_app()) as c:
        assert c.get("/docs").status_code == 403
        assert c.delete("/api/customer/import").status_code == 403
        assert c.post("/api/not-a-real-endpoint").status_code == 403
        assert c.post("/api/forecast/history/save", headers={"Origin": "https://unrelated.invalid"}).status_code == 403
        assert c.post("/api/forecast/history/save", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403


def test_ai_rate_limit_still_applies():
    with client(mock_app()) as c:
        for _ in range(3):
            assert c.post("/api/agent").status_code == 200
        assert c.post("/api/agent").status_code == 429


def test_bundled_synthetic_csv_and_xlsx_are_mappable(monkeypatch):
    monkeypatch.setattr(demo, "_customer_csv", lambda forecast_start=None: FAKE)
    demo.sample_bytes.cache_clear()
    try:
        receipt = demo.sample_bytes("receipts.csv")
        payment = demo.sample_bytes("payments.csv")
        assert b"h1" in receipt and b"h2" not in receipt
        assert b"h2" in payment and b"h1" not in payment
        xlsx = demo.sample_bytes("cash.xlsx")
        assert xlsx[:2] == b"PK"
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(xlsx))
        assert book.active.max_row == 3
        book.close()
        monkeypatch.setenv("RISKPILOT_PUBLIC_INTERACTIVE", "1")
        for filename in ("cash.csv", "receipts.csv", "payments.csv", "cash.xlsx"):
            demo.require_demo_upload(demo.sample_bytes(filename), filename, purpose="mapping")
        demo.require_demo_upload(demo.sample_bytes("cash.csv"), "cash.csv", purpose="customer")
    finally:
        demo.sample_bytes.cache_clear()


def test_arbitrary_upload_rejected_in_public_but_allowed_locally(monkeypatch):
    monkeypatch.setattr(demo, "_customer_csv", lambda forecast_start=None: FAKE)
    demo.sample_bytes.cache_clear()
    try:
        monkeypatch.setenv("RISKPILOT_PUBLIC_INTERACTIVE", "1")
        with pytest.raises(HTTPException) as e:
            demo.require_demo_upload(b"private bank statement", "bank.csv", purpose="mapping")
        assert e.value.status_code == 403
        with pytest.raises(HTTPException):
            demo.require_demo_upload(b"private bank statement", "bank.csv", purpose="customer")
        monkeypatch.delenv("RISKPILOT_PUBLIC_INTERACTIVE", raising=False)
        demo.require_demo_upload(b"private bank statement", "bank.csv", purpose="mapping")
    finally:
        demo.sample_bytes.cache_clear()


def test_missing_public_request_context_fails_closed(monkeypatch):
    monkeypatch.setenv("RISKPILOT_PUBLIC_INTERACTIVE", "1")
    with pytest.raises(HTTPException) as e:
        demo.current_public_owner()
    assert e.value.status_code == 403


def test_history_key_is_visitor_scoped(monkeypatch):
    monkeypatch.setenv("RISKPILOT_PUBLIC_INTERACTIVE", "1")
    t1 = demo.set_public_visitor("visitor-a")
    try:
        a = demo.history_storage_key("same-key")
    finally:
        demo.reset_public_visitor(t1)
    t2 = demo.set_public_visitor("visitor-b")
    try:
        b = demo.history_storage_key("same-key")
    finally:
        demo.reset_public_visitor(t2)
    assert a != b


def test_recently_downloaded_dated_sample_is_accepted(monkeypatch):
    """A generated test file remains eligible after the reviewer downloads it."""
    original = FAKE
    shifted = original.replace(b"2026-10-08", b"2026-10-11").replace(b"2026-10-09", b"2026-10-12")
    monkeypatch.setattr(demo, "_customer_csv", lambda forecast_start=None: shifted if forecast_start else original)
    demo.sample_bytes.cache_clear()
    try:
        response = demo.download_synthetic_sample("cash.csv", forecast_start="2026-10-11")
        assert response.status_code == 200
        assert b"2026-10-11" in response.body
        demo.require_demo_upload(response.body, "RiskPilot_Synthetic_cash.csv", purpose="customer")
        with pytest.raises(HTTPException):
            demo.require_demo_upload(response.body+b"private-change", "cash.csv", purpose="customer")
    finally:
        demo.sample_bytes.cache_clear()
